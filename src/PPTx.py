import numpy as np
import pandas as pd
import openpyxl
from openpyxl.utils import get_column_letter

# Required for PowerPoint chart and presentation manipulations
from pptx import Presentation
from pptx.chart.data import CategoryChartData


# ==========================================
# Excel Functions
# ==========================================

def write_named_range(wb, range_name, values):
    nr = wb.defined_names[range_name]

    sheet_name, cell_range = list(nr.destinations)[0]
    ws = wb[sheet_name]

    start_cell, end_cell = (
        cell_range.split(":")
        if ":" in cell_range
        else (cell_range, cell_range)
    )

    start_row = ws[start_cell].row
    start_col = ws[start_cell].column

    end_row = ws[end_cell].row
    for r in range(start_row, ws.max_row + 1):
        ws.cell(r, start_col).value = None

    # Write new values
    for i, value in enumerate(values):
        ws.cell(start_row + i, start_col, value)

    # Resize named range
    col_letter = get_column_letter(start_col)

    nr.attr_text = (
        f"'{sheet_name}'!"
        f"${col_letter}${start_row}:"
        f"${col_letter}${start_row + len(values) - 1}"
    )


def weighted_score(df, score_col, weight_col="Market Value %"):
    mask = df[score_col].notna()

    if not mask.any():
        return np.nan

    weights = df.loc[mask, weight_col]
    weights = weights / weights.sum()  # normalisation

    return (weights * df.loc[mask, score_col]).sum()


def weighted_score_dataframe(df, score_col, breakdown, weight_col="Market Value %"):
    mask = df[score_col].notna()

    if not mask.any():
        return pd.Series(dtype=float)

    # Normalise weights only on covered securities
    weights = df.loc[mask, weight_col]
    weights = weights / weights.sum()

    weighted_values = weights * df.loc[mask, score_col]
    # return weighted_values.groupby(df.loc[mask, breakdown])


# ==========================================
# PowerPoint Functions
# ==========================================

def update_chart(chart, df, percentage=False, decimals=1):
    chart_data = CategoryChartData()

    chart_data.categories = df.index.astype(str).tolist()

    for col in df.columns:
        chart_data.add_series(
            str(col),
            df[col].tolist()
        )

    chart.replace_data(chart_data)

    # Format d'affichage
    fmt = f"0.{'0' * decimals}%" if decimals > 0 else "0%"

    if percentage:
        try:
            # Axe Y
            chart.value_axis.tick_labels.number_format = fmt
        except Exception:
            pass

        # Data Labels
        try:
            for series in chart.series:
                if series.has_data_labels:
                    series.data_labels.number_format = fmt
        except Exception:
            pass


def get_chart(prs, chart_name):
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_chart and shape.name == chart_name:
                return shape

    raise ValueError(f"Chart '{chart_name}' not found")


def update_text(shape, new_text):
    if not shape.has_text_frame:
        return

    paragraph = shape.text_frame.paragraphs[0]

    if paragraph.runs:
        paragraph.runs[0].text = new_text

        # Remove extra runs if present
        while len(paragraph.runs) > 1:
            paragraph._p.remove(paragraph.runs[-1]._r)
    else:
        paragraph.text = new_text


def update_table(table, df):
    # Write index
    for r in range(len(df)):
        cell = table.cell(r + 1, 0)
        value = str(df.index[r])

        p = cell.text_frame.paragraphs[0]
        if p.runs:
            p.runs[0].text = value
        else:
            p.text = value

    # Write data
    for r in range(len(df)):
        for c in range(len(df.columns)):
            cell = table.cell(r + 1, c + 1)
            value = str(df.iloc[r, c])

            p = cell.text_frame.paragraphs[0]
            if p.runs:
                p.runs[0].text = value
            else:
                p.text = value


def update_chart_title(chart, new_title):
    chart.has_title = True

    p = chart.chart_title.text_frame.paragraphs[0]

    if p.runs:
        p.runs[0].text = str(new_title)

        while len(p.runs) > 1:
            p._p.remove(p.runs[-1]._r)
    else:
        p.text = str(new_title)

